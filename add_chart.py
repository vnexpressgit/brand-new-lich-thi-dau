import sys
import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList

def add_chart():
    file_path = 'data-dlbd.xlsx'
    print(f"Loading {file_path}...")
    wb = openpyxl.load_workbook(file_path)
    
    ws = wb['Thống kê']
    print(f"Loaded sheet: {ws.title}")

    # Clear existing charts if any
    ws._charts.clear()

    # Find the header row for Section III
    header_row = None
    for r in range(80, 95):
        val = ws.cell(row=r, column=2).value
        if val == 'Tháng':
            header_row = r
            break
            
    if not header_row:
        raise ValueError("Could not find 'Tháng' header in rows 80-95")

    first_data_row = header_row + 1
    last_data_row = header_row + 21  # 21 months of data

    print(f"Header row: {header_row}, Data rows: {first_data_row} to {last_data_row}")
    print(f"First month: {ws.cell(row=first_data_row, column=2).value}, Last month: {ws.cell(row=last_data_row, column=2).value}")

    # Create BarChart (column chart)
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "BIỂU ĐỒ BIẾN ĐỘNG TỔNG PAGEVIEW THEO TỪNG THÁNG (01/2025 - 09/2026)"
    chart.y_axis.title = "Lượt xem (Pageview)"
    chart.x_axis.title = "Tháng"
    chart.legend = None

    # Data range: Col 3 (Tổng Pageview), rows header_row to last_data_row
    data = Reference(ws, min_col=3, min_row=header_row, max_row=last_data_row)
    # Categories: Col 2 (Tháng), rows first_data_row to last_data_row
    cats = Reference(ws, min_col=2, min_row=first_data_row, max_row=last_data_row)

    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)

    # Style series color: Elegant Navy Blue #1F4E79
    if chart.series:
        chart.series[0].graphicalProperties.solidFill = "1F4E79"

    # Show data labels on top of each bar
    chart.dataLabels = DataLabelList()
    chart.dataLabels.showVal = True

    # Chart size: width ~25cm (~950px), height ~13.5cm (~510px)
    chart.width = 25
    chart.height = 13.5

    # Position chart starting at cell A110 (3 rows below row 107 total row)
    chart_cell = f"A{last_data_row + 4}"
    print(f"Adding chart at cell {chart_cell}...")
    ws.add_chart(chart, chart_cell)

    print("Saving workbook...")
    wb.save(file_path)
    print("SUCCESS: Chart added and saved to data-dlbd.xlsx!")

if __name__ == '__main__':
    add_chart()
